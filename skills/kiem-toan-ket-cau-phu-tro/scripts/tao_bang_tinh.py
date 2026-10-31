#!/usr/bin/env python3
"""Tạo khung bảng tính kiểm toán kết cấu phụ trợ theo references/bang-tinh-excel.md.

Ví dụ:
  python tao_bang_tinh.py --out KT_da-giao.xlsx --cong-trinh "Cầu X" \
      --hang-muc "Đà giáo nhịp 3" --he cau --nhom VK DG NM

Sau khi điền số liệu, chạy recalc.py của skill xlsx để tính lại công thức.
"""
import argparse, json, os, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")

FONT = "Arial"
F_NORMAL = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=13, bold=True)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
FILL_INPUT = PatternFill("solid", fgColor="FFF9DB")
FILL_HEAD = PatternFill("solid", fgColor="D9E2F3")
FILL_EX = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

TEN_NHOM = {"VK": "Ván khuôn", "DG": "Đà giáo, cột chống, trụ tạm", "NM": "Nền, móng tạm",
            "VV": "Vòng vây, hố đào, văng chống, neo", "TB": "Xe đúc, MSS, giá lao, cẩu",
            "CT": "Cầu tạm, sàn đạo, sàn thao tác", "TH": "Công trình thủy tạm, hệ nổi"}
DANG_LIST = ["S≤R", "S·K≤R", "δ≤[δ]", "KT", "Tải/nhập liệu"]
VAI_TRO = {
    "cau": {"chính": ["N01"], "bắt buộc": ["N05"], "khi dẫn chiếu": ["N06", "N07", "N10", "N11", "N14", "N15", "N16"]},
    "dan-dung": {"chính": ["N02", "N03", "N04"], "bắt buộc": ["N05"], "khi dẫn chiếu": ["N07", "N08", "N09", "N13", "N14", "N16", "N17"]},
    "cang": {"chính": ["N12", "N01"], "bắt buộc": ["N05"], "khi dẫn chiếu": ["N04", "N07", "N10", "N11", "N14", "N15"]},
}


def dang_tach(dang):
    """'S ≤ R; δ ≤ [δ]' -> ['S≤R','δ≤[δ]']"""
    out = []
    for part in re.split(r";|\+", dang):
        p = part.replace(" ", "").strip()
        if not p:
            continue
        if p.startswith("S·K"):
            out.append("S·K≤R")
        elif p.startswith("S≤R"):
            out.append("S≤R")
        elif p.startswith("δ"):
            out.append("δ≤[δ]")
        elif p.startswith("KT"):
            out.append("KT")
        else:
            out.append("Tải/nhập liệu")
    seen = []
    for d in out:
        if d not in seen:
            seen.append(d)
    return seen or ["S≤R"]


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.border, c.alignment = F_BOLD, FILL_HEAD, BOX, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE


def inp(c, value=None):
    if value is not None:
        c.value = value
    c.font, c.fill, c.border, c.alignment = F_INPUT, FILL_INPUT, BOX, WRAP


def norm(c, value=None, bold=False):
    if value is not None:
        c.value = value
    c.font, c.border, c.alignment = (F_BOLD if bold else F_NORMAL), BOX, WRAP


def sheet_kiem_tra(wb, nhom, rows, he):
    ws = wb.create_sheet(f"KT_{nhom}")
    title(ws, f"Kiểm tra nhóm {nhom} - {TEN_NHOM[nhom]}",
          "Ô vàng chữ xanh là ô nhập. Cột 'Nguồn công thức' để trống thì kết luận là CHƯA CÓ CĂN CỨ. "
          "Cột 'Điều khoản gợi ý' lấy từ ma trận, phải đối chiếu bản gốc rồi mới chép sang cột nguồn.")
    cols = ["Mã", "Trạng thái kiểm tra", "Giai đoạn", "Dạng điều kiện", "S (tác động)", "R (khả năng / giới hạn)",
            "K (hệ số an toàn / tin cậy)", "Đơn vị", "Hệ số sử dụng u", "Kết luận", "Nguồn công thức",
            "Nguồn R và K", "Điều khoản gợi ý (ma trận)", "Ghi chú"]
    header(ws, 4, cols, [7, 38, 16, 12, 12, 14, 12, 9, 12, 20, 22, 22, 40, 28])
    dv = DataValidation(type="list", formula1='"' + ",".join(DANG_LIST) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    r = 5
    for item in rows:
        goi_y = item["cau"] if he == "cau" else item["khac"] or item["cau"]
        for d in dang_tach(item["dang"]):
            ten = item["trang_thai"] + (" - biến dạng" if d == "δ≤[δ]" and "δ" not in item["trang_thai"] else "")
            norm(ws.cell(r, 1), item["ma"])
            norm(ws.cell(r, 2), ten)
            inp(ws.cell(r, 3))
            inp(ws.cell(r, 4), d); dv.add(ws.cell(r, 4))
            for col in (5, 6, 7, 8):
                inp(ws.cell(r, col))
            norm(ws.cell(r, 9), f'=IF(OR(D{r}="KT",D{r}="Tải/nhập liệu",E{r}="",F{r}=""),"",'
                                f'IFERROR(IF(D{r}="S·K≤R",E{r}*IF(G{r}="",1,G{r})/F{r},E{r}/F{r}),"LỖI R=0"))')
            ws.cell(r, 9).number_format = "0.000"
            norm(ws.cell(r, 10), f'=IF(D{r}="Tải/nhập liệu","-",IF(D{r}="KT","KIỂM CẤU TẠO",'
                                 f'IF(K{r}="","CHƯA CÓ CĂN CỨ",IF(I{r}="","CHƯA TÍNH",'
                                 f'IF(I{r}="LỖI R=0","CHƯA TÍNH",IF(I{r}<=1,"ĐẠT","KHÔNG ĐẠT"))))))')
            inp(ws.cell(r, 11)); inp(ws.cell(r, 12))
            norm(ws.cell(r, 13), goi_y)
            ws.cell(r, 13).font = F_NOTE
            inp(ws.cell(r, 14))
            r += 1
    last = r - 1
    rng = f"J5:J{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=['J5="ĐẠT"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['J5="KHÔNG ĐẠT"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['OR(J5="CHƯA CÓ CĂN CỨ",J5="CHƯA TÍNH")'],
                                                   fill=PatternFill("solid", fgColor="FFEB9C")))
    return ws, last


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--cong-trinh", default="")
    ap.add_argument("--hang-muc", default="")
    ap.add_argument("--nguoi-lap", default="")
    ap.add_argument("--he", choices=["cau", "dan-dung", "cang"], default="cau")
    ap.add_argument("--nhom", nargs="+", default=["VK", "DG", "NM"], choices=list(TEN_NHOM))
    ap.add_argument("--khoa", action="store_true", help="Khóa ô công thức (không đặt mật khẩu)")
    a = ap.parse_args()

    mt = json.load(open(os.path.join(ASSETS, "ma-tran.json"), encoding="utf-8"))
    ng = json.load(open(os.path.join(ASSETS, "so-nguon.json"), encoding="utf-8"))

    wb = Workbook()
    th = wb.active; th.title = "Tong_hop"

    # Thong_tin
    ws = wb.create_sheet("Thong_tin")
    title(ws, "Thông tin hồ sơ kiểm toán")
    ws.column_dimensions["A"].width = 32; ws.column_dimensions["B"].width = 60
    info = [("Công trình", a.cong_trinh), ("Hạng mục kết cấu phụ trợ", a.hang_muc),
            ("Hệ tiêu chuẩn chính", {"cau": "Cầu - TCVN 11815:2017 [N01]", "dan-dung": "Dân dụng - TCVN 4453 [N02] + TCVN 2737:2023 [N03] + TCVN 5575:2024 [N04]", "cang": "Cảng - TCVN 11820 [N12] + TCVN 11815 [N01]"}[a.he]),
            ("Phiên bản", "V0"), ("Ngày lập", datetime.date.today().isoformat()), ("Người lập", a.nguoi_lap),
            ("Người soát", ""), ("Hồ sơ biện pháp dẫn chiếu", ""), ("Thư mục lưu hồ sơ", "")]
    for i, (k, v) in enumerate(info, start=4):
        norm(ws.cell(i, 1), k, bold=True); inp(ws.cell(i, 2), v)

    # Nguon
    ws = wb.create_sheet("Nguon")
    title(ws, "Sổ nguồn", "Vai trò gợi ý theo hệ đã chọn. Cột 'Bản gốc đã nhận' và 'Hiệu lực' do người dùng xác nhận.")
    header(ws, 4, ["Mã", "Số hiệu", "Tên / nội dung", "Dùng cho", "Vai trò trong bài", "Bản gốc đã nhận", "Hiệu lực (người dùng xác nhận)", "Xác minh khi lập skill"],
           [7, 26, 40, 50, 16, 14, 20, 45])
    vt = VAI_TRO[a.he]
    for i, n in enumerate(ng, start=5):
        role = next((k for k, v in vt.items() if n["ma"] in v), "đối chiếu" if n["ma"].startswith("Q") else ("dự án" if n["ma"].startswith("D") else ""))
        for j, key in enumerate(["ma", "so_hieu", "ten", "dung_cho"], start=1):
            norm(ws.cell(i, j), n[key])
        inp(ws.cell(i, 5), role); inp(ws.cell(i, 6)); inp(ws.cell(i, 7))
        norm(ws.cell(i, 8), n["xac_minh"]); ws.cell(i, 8).font = F_NOTE
    r0 = len(ng) + 7
    ws.cell(r0, 1, "Chỉ mục trị số đã trích từ bản gốc").font = F_BOLD
    header_row = r0 + 1
    for j, (t, w) in enumerate(zip(["Mã", "Điều / CT / Bảng", "Trang", "Nội dung", "Trị số / biểu thức", "Đơn vị", "Điều kiện áp dụng"], [7, 26, 40, 50, 16, 14, 20]), start=1):
        c = ws.cell(header_row, j, t); c.font, c.fill, c.border = F_BOLD, FILL_HEAD, BOX
    for j in range(1, 8):
        inp(ws.cell(header_row + 1, j))

    # Dau_vao
    ws = wb.create_sheet("Dau_vao")
    title(ws, "Dữ liệu đầu vào", "Mỗi dòng một đại lượng; công thức ở sheet khác tham chiếu tới ô cột D. Dòng xám là VÍ DỤ ĐỊNH DẠNG, xóa khi dùng.")
    header(ws, 4, ["Mã", "Đại lượng", "Ký hiệu", "Giá trị", "Đơn vị", "Nguồn số liệu", "Giả định?", "Ghi chú"], [8, 36, 10, 12, 10, 36, 10, 30])
    ex = ["DV-01", "Chiều cao đổ bê tông một đợt (ví dụ)", "H", 3.0, "m", "Bản vẽ biện pháp BPTC-xx (ví dụ)", "không", "VÍ DỤ - xóa"]
    for j, v in enumerate(ex, start=1):
        c = ws.cell(5, j, v); c.font, c.fill, c.border = F_INPUT, FILL_EX, BOX
    for rr in range(6, 30):
        for j in range(1, 9):
            inp(ws.cell(rr, j))

    # Tai_trong
    ws = wb.create_sheet("Tai_trong")
    title(ws, "Tải trọng", "Lập đủ theo danh mục của hệ đã chọn (N01: Bảng 5.1). Tải không xuất hiện ghi 'không xuất hiện' và lý do.")
    header(ws, 4, ["Mã tải", "Loại tải", "Giá trị tiêu chuẩn", "Đơn vị", "Hệ số lớn", "Hệ số nhỏ", "Xung kích (1+μ)", "Giai đoạn xuất hiện", "Nguồn giá trị", "Nguồn hệ số", "Tính toán (lớn)", "Tính toán (nhỏ)"],
           [8, 34, 14, 9, 10, 10, 12, 22, 22, 22, 14, 14])
    for rr in range(5, 30):
        for j in range(1, 11):
            inp(ws.cell(rr, j))
        norm(ws.cell(rr, 11), f'=IF(C{rr}="","",C{rr}*IF(E{rr}="",1,E{rr})*IF(G{rr}="",1,G{rr}))')
        norm(ws.cell(rr, 12), f'=IF(C{rr}="","",C{rr}*IF(F{rr}="",1,F{rr}))')

    # To_hop
    ws = wb.create_sheet("To_hop")
    title(ws, "Tổ hợp tải trọng", "Mỗi tổ hợp một dòng: liệt kê mã tải và hệ số dùng, nguồn quy định tổ hợp. Đánh dấu tổ hợp khống chế.")
    header(ws, 4, ["Tổ hợp", "Giai đoạn", "Mục đích (TTGH1 / TTGH2 / ổn định)", "Các tải và hệ số", "Hệ số tổ hợp", "Nguồn", "Khống chế cho"], [10, 20, 22, 50, 12, 24, 24])
    for rr in range(5, 20):
        for j in range(1, 8):
            inp(ws.cell(rr, j))

    # Kiem tra
    kt_info = []
    for nhom in a.nhom:
        rows = [m for m in mt if m["nhom"] == nhom]
        sh, last = sheet_kiem_tra(wb, nhom, rows, a.he)
        kt_info.append((sh.title, last))

    # Gia_dinh, Chua_co_can_cu, Nhat_ky
    ws = wb.create_sheet("Gia_dinh")
    title(ws, "Giả định chưa xác nhận")
    header(ws, 4, ["STT", "Giả định", "Ảnh hưởng tới kiểm tra", "Người xác nhận", "Tình trạng", "Ghi chú"], [6, 50, 30, 18, 16, 30])
    dv2 = DataValidation(type="list", formula1='"Chưa xác nhận,Đã xác nhận,Bác bỏ"', allow_blank=True); ws.add_data_validation(dv2)
    for rr in range(5, 30):
        for j in range(1, 7):
            inp(ws.cell(rr, j))
        dv2.add(ws.cell(rr, 5))
    ws = wb.create_sheet("Chua_co_can_cu")
    title(ws, "Nội dung chưa có căn cứ")
    header(ws, 4, ["STT", "Mã kiểm tra", "Nội dung cần", "Tài liệu / điều khoản cần bổ sung", "Người cung cấp", "Tình trạng"], [6, 12, 40, 40, 18, 16])
    for rr in range(5, 25):
        for j in range(1, 7):
            inp(ws.cell(rr, j))
    ws = wb.create_sheet("Nhat_ky")
    title(ws, "Nhật ký sửa đổi")
    header(ws, 4, ["Phiên bản", "Ngày", "Người sửa", "Nội dung sửa"], [10, 12, 18, 70])
    for j, v in enumerate(["V0", datetime.date.today().isoformat(), a.nguoi_lap, "Tạo khung"], start=1):
        inp(ws.cell(5, j), v)

    # Tong_hop
    title(th, "TỔNG HỢP KIỂM TOÁN KẾT CẤU PHỤ TRỢ")
    th["A2"] = "=Thong_tin!B4&\" - \"&Thong_tin!B5"; th["A2"].font = F_BOLD
    th.column_dimensions["A"].width = 30
    for col, w in zip("BCDEFG", [14, 16, 16, 18, 16, 16]):
        th.column_dimensions[col].width = w
    header(th, 4, ["Sheet kiểm tra", "u lớn nhất", "Mã khống chế", "Số KHÔNG ĐẠT", "Số CHƯA CÓ CĂN CỨ", "Số CHƯA TÍNH", "Số ĐẠT"], [30, 14, 16, 16, 18, 16, 16])
    r = 5
    for name, last in kt_info:
        norm(th.cell(r, 1), name)
        norm(th.cell(r, 2), f"=IF(COUNT({name}!I5:I{last})=0,\"\",MAX({name}!I5:I{last}))"); th.cell(r, 2).number_format = "0.000"
        norm(th.cell(r, 3), f"=IF(B{r}=\"\",\"\",IFERROR(INDEX({name}!A5:A{last},MATCH(B{r},{name}!I5:I{last},0)),\"\"))")
        for col, key in zip((4, 5, 6, 7), ("KHÔNG ĐẠT", "CHƯA CÓ CĂN CỨ", "CHƯA TÍNH", "ĐẠT")):
            norm(th.cell(r, col), f'=COUNTIF({name}!J5:J{last},"{key}")')
        r += 1
    tot = r
    norm(th.cell(tot, 1), "Tổng", bold=True)
    norm(th.cell(tot, 2), f'=IF(COUNT(B5:B{tot-1})=0,"",MAX(B5:B{tot-1}))'); th.cell(tot, 2).number_format = "0.000"
    for col in (4, 5, 6, 7):
        L = get_column_letter(col)
        norm(th.cell(tot, col), f"=SUM({L}5:{L}{tot-1})")
    r = tot + 2
    norm(th.cell(r, 1), "Giả định chưa xác nhận", bold=True)
    norm(th.cell(r, 2), '=COUNTIF(Gia_dinh!E5:E29,"Chưa xác nhận")')
    r += 1
    norm(th.cell(r, 1), "KẾT LUẬN", bold=True)
    th.cell(r, 2, f'=IF(D{tot}>0,"KHÔNG ĐẠT",IF(OR(E{tot}>0,F{tot}>0,B{r-1}>0,G{tot}=0),"CHƯA ĐỦ CƠ SỞ KẾT LUẬN","ĐẠT"))')
    th.cell(r, 2).font = Font(name=FONT, size=12, bold=True)
    th.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    th.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f'B{r}="ĐẠT"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    th.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f'B{r}="KHÔNG ĐẠT"'], fill=PatternFill("solid", fgColor="FFC7CE")))
    th.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f'B{r}="CHƯA ĐỦ CƠ SỞ KẾT LUẬN"'], fill=PatternFill("solid", fgColor="FFEB9C")))
    r += 2
    norm(th.cell(r, 1), "Điều kiện áp dụng của kết luận", bold=True)
    for k in ["Tải lớn nhất cho phép", "Trình tự thi công", "Tốc độ đổ bê tông lớn nhất", "Mực nước lớn nhất / nhỏ nhất", "Gió giới hạn khi làm việc", "Ngưỡng quan trắc"]:
        r += 1
        norm(th.cell(r, 1), k); inp(th.cell(r, 2)); th.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
    r += 2
    th.cell(r, 1, "Kết quả do người có chuyên môn kiểm tra và ký duyệt. Khuyến nghị thẩm tra độc lập với đà giáo cao - nhịp lớn, xe đúc, MSS, giá lao, hố đào sâu, vòng vây dưới nước.").font = F_NOTE
    r += 1
    th.cell(r, 1, "Chú giải: ô vàng chữ xanh là ô nhập; ô chữ đen là công thức; kết luận chỉ 'ĐẠT' khi mọi dòng đã tính, có nguồn và không còn giả định treo.").font = F_NOTE

    if a.khoa:
        for sh in wb.worksheets:
            sh.protection.sheet = True
            for row in sh.iter_rows():
                for c in row:
                    if c.fill == FILL_INPUT:
                        c.protection = c.protection.copy(locked=False)

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    wb.save(a.out)
    print("Đã tạo:", a.out, "| sheet kiểm tra:", ", ".join(n for n, _ in kt_info))


if __name__ == "__main__":
    main()
