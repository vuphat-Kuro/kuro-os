#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docgen — Trộn dữ liệu Excel vào mẫu Word để sinh hồ sơ hàng loạt.

Dùng cho: biên bản nghiệm thu công việc, phiếu yêu cầu nghiệm thu, phiếu
kiểm tra, tờ trình, công văn — mọi loại văn bản chỉ khác nhau vài trường
nhưng phải in ra hàng trăm bản.

Mẫu Word dùng chỗ trống dạng {{Tên cột}} — trùng đúng tên cột trong Excel.
Chỗ trống đặt được ở thân văn bản, trong bảng, ở header và footer.

Các lệnh:
    tao        Sinh hàng loạt file Word từ Excel + mẫu
    cho-trong  Liệt kê mọi chỗ trống có trong mẫu (chạy trước khi tạo)
    kiem-tra   Đối chiếu mẫu với Excel: thiếu cột nào, thừa chỗ trống nào

Ví dụ:
    python docgen.py cho-trong --mau mau-bien-ban.docx
    python docgen.py kiem-tra  --mau mau-bien-ban.docx --du-lieu ds.xlsx
    python docgen.py tao --mau mau-bien-ban.docx --du-lieu ds.xlsx \\
        --ra-thu-muc bien-ban/ --ten "BB-{{So hieu}}-{{Hang muc}}.docx"
"""

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document

CHO_TRONG = re.compile(r"\{\{\s*([^}]+?)\s*\}\}")

THANG_VN = ["một", "hai", "ba", "tư", "năm", "sáu", "bảy", "tám",
            "chín", "mười", "mười một", "mười hai"]


# ----------------------------------------------------------------------------
# Đọc/ghi chỗ trống trong docx
# ----------------------------------------------------------------------------

def moi_doan(doc):
    """Duyệt mọi đoạn văn trong tài liệu: thân, bảng (kể cả bảng lồng),
    header và footer của mọi section."""
    def tu_bang(bang):
        for hang in bang.rows:
            for o in hang.cells:
                for p in o.paragraphs:
                    yield p
                for b2 in o.tables:
                    yield from tu_bang(b2)

    for p in doc.paragraphs:
        yield p
    for b in doc.tables:
        yield from tu_bang(b)
    for sec in doc.sections:
        for phan in (sec.header, sec.footer,
                     sec.first_page_header, sec.first_page_footer,
                     sec.even_page_header, sec.even_page_footer):
            if phan is None:
                continue
            for p in phan.paragraphs:
                yield p
            for b in phan.tables:
                yield from tu_bang(b)


def thay_trong_doan(doan, gia_tri):
    """Thay chỗ trống trong một đoạn, giữ nguyên định dạng.

    Word hay cắt '{{Ten}}' thành nhiều run rời rạc (do gõ, do sửa chính tả),
    nên thay từng run sẽ trượt. Cách làm: dồn text về run đầu, xóa text các
    run sau, rồi thay trên chuỗi đã dồn. Định dạng của run đầu được giữ.
    """
    van_ban = "".join(r.text for r in doan.runs)
    if "{{" not in van_ban:
        return 0

    def doi(m):
        khoa = m.group(1).strip()
        v = gia_tri.get(khoa)
        if v is None:
            v = gia_tri.get(khoa.lower())
        return "" if v is None else str(v)

    moi = CHO_TRONG.sub(doi, van_ban)
    if moi == van_ban:
        return 0

    if doan.runs:
        doan.runs[0].text = moi
        for r in doan.runs[1:]:
            r.text = ""
    else:
        doan.text = moi
    return len(CHO_TRONG.findall(van_ban))


def lay_cho_trong(duong_dan):
    doc = Document(duong_dan)
    ds = []
    for p in moi_doan(doc):
        van_ban = "".join(r.text for r in p.runs) or p.text
        for m in CHO_TRONG.finditer(van_ban):
            ten = m.group(1).strip()
            if ten not in ds:
                ds.append(ten)
    return ds


# ----------------------------------------------------------------------------
# Trường tự sinh
# ----------------------------------------------------------------------------

def truong_tu_sinh(so_thu_tu):
    hom_nay = datetime.now()
    return {
        "STT tự động": so_thu_tu,
        "Ngày": hom_nay.strftime("%d"),
        "Tháng": hom_nay.strftime("%m"),
        "Năm": hom_nay.strftime("%Y"),
        "Ngày tháng": hom_nay.strftime("%d/%m/%Y"),
        "Ngày tháng đầy đủ": (
            f"ngày {hom_nay.day:02d} tháng {THANG_VN[hom_nay.month - 1]} "
            f"năm {hom_nay.year}"
        ),
    }


def dinh_dang_gia_tri(v, dinh_dang_so=True):
    """Số trong hồ sơ Việt Nam viết 1.234,56 — không phải 1234.56."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, pd.Timestamp) or isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if dinh_dang_so and isinstance(v, (int, float)) and not isinstance(v, bool):
        if float(v).is_integer():
            return f"{int(v):,}".replace(",", ".")
        # 125.5 -> "125,5" (bỏ số 0 thừa), 1234.567 -> "1.234,567"
        t = f"{v:,.3f}".rstrip("0").rstrip(".")
        return t.replace(",", "\x00").replace(".", ",").replace("\x00", ".")
    return str(v)


def ten_an_toan(s):
    s = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", str(s))
    return re.sub(r"\s+", " ", s).strip()[:120] or "khong-ten"


# ----------------------------------------------------------------------------
# Lệnh
# ----------------------------------------------------------------------------

def lenh_cho_trong(a):
    ds = lay_cho_trong(a.mau)
    if not ds:
        print("⚠️  Mẫu không có chỗ trống nào dạng {{...}}.")
        print("    Mở file Word, gõ {{Tên cột}} vào chỗ cần điền rồi lưu lại.")
        return
    print(f"📋 Mẫu '{Path(a.mau).name}' có {len(ds)} chỗ trống:\n")
    for t in ds:
        print(f"   {{{{{t}}}}}")
    print("\n→ Đặt tên cột trong Excel trùng đúng các tên trên.")
    print("   Trường tự sinh dùng được ngay, không cần cột trong Excel:")
    for k in truong_tu_sinh(1):
        print(f"   {{{{{k}}}}}")


def lenh_kiem_tra(a):
    ds = lay_cho_trong(a.mau)
    df = pd.read_excel(a.du_lieu, sheet_name=a.sheet)
    cot = [str(c).strip() for c in df.columns]
    tu_sinh = set(truong_tu_sinh(1))

    thieu = [t for t in ds if t not in cot and t not in tu_sinh]
    thua = [c for c in cot if c not in ds]

    print(f"Mẫu: {Path(a.mau).name} — {len(ds)} chỗ trống")
    print(f"Dữ liệu: {Path(a.du_lieu).name} — {len(df)} dòng, {len(cot)} cột\n")

    if thieu:
        print(f"🔴 {len(thieu)} chỗ trống KHÔNG có cột tương ứng "
              f"(sẽ bị thay bằng rỗng):")
        for t in thieu:
            gan = [c for c in cot if t.lower().replace(" ", "") in c.lower().replace(" ", "")
                   or c.lower().replace(" ", "") in t.lower().replace(" ", "")]
            print(f"   {{{{{t}}}}}" + (f"  → có phải ý bạn là cột '{gan[0]}'?" if gan else ""))
    else:
        print("✅ Mọi chỗ trống đều có cột tương ứng.")

    if thua:
        print(f"\n⚪ {len(thua)} cột trong Excel không dùng đến: {', '.join(thua[:12])}")

    if thieu:
        sys.exit(2)


def lenh_tao(a):
    doc_kieu = {c.strip(): str for c in a.cot_chu.split(",") if c.strip()}
    df = pd.read_excel(a.du_lieu, sheet_name=a.sheet, dtype=doc_kieu or None)
    df.columns = [str(c).strip() for c in df.columns]
    if a.loc:
        cot, gia_tri = a.loc.split("=", 1)
        df = df[df[cot.strip()].astype(str).str.strip() == gia_tri.strip()]
    if a.tu_dong is not None:
        df = df.head(a.tu_dong)
    if df.empty:
        raise SystemExit("❌ Không có dòng dữ liệu nào để tạo.")

    thu_muc = Path(a.ra_thu_muc)
    thu_muc.mkdir(parents=True, exist_ok=True)
    ds_cho_trong = lay_cho_trong(a.mau)
    da_tao, ten_da_dung = [], {}

    for i, (_, hang) in enumerate(df.iterrows(), start=1):
        gia_tri = {k: dinh_dang_gia_tri(v, not a.giu_so_tho)
                   for k, v in hang.to_dict().items()}
        gia_tri.update({k: str(v) for k, v in truong_tu_sinh(i).items()})

        doc = Document(a.mau)
        for p in moi_doan(doc):
            thay_trong_doan(p, gia_tri)

        ten = a.ten
        for m in CHO_TRONG.finditer(a.ten):
            ten = ten.replace(m.group(0), ten_an_toan(gia_tri.get(m.group(1).strip(), "")))
        if not ten.lower().endswith(".docx"):
            ten += ".docx"
        # tránh ghi đè khi hai dòng ra cùng tên
        if ten in ten_da_dung:
            ten_da_dung[ten] += 1
            ten = f"{Path(ten).stem}_{ten_da_dung[ten]}.docx"
        else:
            ten_da_dung[ten] = 1

        dich = thu_muc / ten
        doc.save(dich)
        da_tao.append({"STT": i, "Tên file": ten,
                       **{k: gia_tri.get(k, "") for k in ds_cho_trong[:5]}})

    print(f"✅ Đã tạo {len(da_tao)} file trong {thu_muc}")
    for r in da_tao[:5]:
        print(f"   • {r['Tên file']}")
    if len(da_tao) > 5:
        print(f"   ... và {len(da_tao) - 5} file nữa")

    if a.nhat_ky:
        nk = Path(a.nhat_ky)
        nk.parent.mkdir(parents=True, exist_ok=True)
        pd.DataFrame(da_tao).to_excel(nk, index=False)
        print(f"   Nhật ký: {nk}")

    if a.pdf:
        chuyen_pdf(thu_muc, [thu_muc / r["Tên file"] for r in da_tao])


def chuyen_pdf(thu_muc, files):
    """Chuyển sang PDF. Ưu tiên LibreOffice (có trên mọi hệ), Windows thì
    dùng Word qua COM nếu đã cài."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        print(f"⏳ Đang chuyển {len(files)} file sang PDF bằng LibreOffice...")
        subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                        "--outdir", str(thu_muc)] + [str(f) for f in files],
                       check=False, capture_output=True)
        print(f"✅ PDF nằm cùng thư mục {thu_muc}")
        return
    try:  # Windows + Microsoft Word
        import win32com.client  # type: ignore
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        for f in files:
            d = word.Documents.Open(str(Path(f).resolve()))
            d.SaveAs(str(Path(f).with_suffix(".pdf").resolve()), FileFormat=17)
            d.Close()
        word.Quit()
        print(f"✅ Đã xuất PDF bằng Microsoft Word")
    except Exception as e:
        print(f"⚠️  Không chuyển được PDF ({e}).")
        print("    Cách 1: cài LibreOffice rồi chạy lại.")
        print("    Cách 2 (Windows có Word): pip install pywin32")


def main():
    p = argparse.ArgumentParser(
        description="docgen — trộn Excel vào mẫu Word, sinh hồ sơ hàng loạt",
        formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    sub = p.add_subparsers(dest="lenh", required=True)

    s = sub.add_parser("cho-trong", help="Liệt kê chỗ trống trong mẫu")
    s.add_argument("--mau", required=True)
    s.set_defaults(func=lenh_cho_trong)

    s = sub.add_parser("kiem-tra", help="Đối chiếu mẫu với file dữ liệu")
    s.add_argument("--mau", required=True)
    s.add_argument("--du-lieu", required=True)
    s.add_argument("--sheet", default=0)
    s.set_defaults(func=lenh_kiem_tra)

    s = sub.add_parser("tao", help="Sinh hàng loạt file Word")
    s.add_argument("--mau", required=True, help="File .docx mẫu")
    s.add_argument("--du-lieu", required=True, help="File .xlsx dữ liệu, mỗi dòng một văn bản")
    s.add_argument("--sheet", default=0)
    s.add_argument("--ra-thu-muc", required=True)
    s.add_argument("--ten", default="{{STT tự động}}.docx",
                   help='Mẫu tên file, ví dụ "BB-{{So hieu}}-{{Hang muc}}.docx"')
    s.add_argument("--loc", default="", help='Chỉ tạo dòng khớp, ví dụ "Hạng mục=Trụ P1"')
    s.add_argument("--tu-dong", type=int, default=None, help="Chỉ tạo N dòng đầu (chạy thử)")
    s.add_argument("--cot-chu", default="",
                   help='Cột giữ nguyên dạng chữ, ví dụ "So hieu" để 01 không thành 1')
    s.add_argument("--giu-so-tho", action="store_true",
                   help="Không định dạng số kiểu 1.234,56")
    s.add_argument("--pdf", action="store_true", help="Xuất thêm bản PDF")
    s.add_argument("--nhat-ky", default="", help="Ghi bảng nhật ký các file đã tạo")
    s.set_defaults(func=lenh_tao)

    a = p.parse_args()
    if isinstance(getattr(a, "sheet", 0), str) and str(a.sheet).isdigit():
        a.sheet = int(a.sheet)
    a.func(a)


if __name__ == "__main__":
    main()
